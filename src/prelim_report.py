"""Preliminary review register - the Stage 1 output that goes to Archer.

The organisation's preliminary review is a table sent to stakeholders:
Domain | Threat | Risk | Risk rating | Cyber recommendation. Stakeholders
build against it, and the final (as-built) review checks each row again.
So every row gets a stable ID (REC-001…) and acceptance criteria - what
the as-built evidence must show - kept with the register so Stage 3 can
verify the same rows.

Rows come only from items a human confirmed: findings the architect
accepted, and threats that are accepted or come from rules over the
approved DFD (unless disputed). The "Risk" statement (business consequence,
not the defect) and acceptance criteria are drafted by the local model and
edited by the architect before export.

The Excel export is deliberately plain for Archer data import: one sheet,
one header row with configurable field names, one row per record, no merged
cells or formulas, rating values from a configurable list.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from .dfd import DFD, slug

DEFAULT_COLUMNS = {
    "id": "ID", "domain": "Domain", "threat": "Threat", "risk": "Risk",
    "rating": "Risk Rating", "recommendation": "Cyber Recommendation",
    "acceptance": "Acceptance Criteria",
}
DEFAULT_RATINGS = {"CRITICAL": "Critical", "HIGH": "High", "MEDIUM": "Medium", "LOW": "Low"}
DEFAULT_DOMAINS = {"network": "Network", "application": "Application",
                   "security": "Cyber Security", "cloud_data": "Cloud & Data",
                   "general": "General"}
# Threats carry no domain; derive it from the DFD element they target.
KIND_DOMAIN = {"network": "network", "device": "network", "datastore": "cloud_data",
               "vector_db": "cloud_data", "identity": "security", "external_party": "security",
               "user_group": "security", "service": "application", "pipeline": "application",
               "llm": "application", "agent": "application", "other": "application"}
ORDER = ["CRITICAL", "HIGH", "MEDIUM", "LOW"]
DEFAULT_ROOT = Path("data/prelim")


@dataclass
class PrelimRow:
    id: str
    domain: str                  # internal key: network / application / security / cloud_data
    threat: str
    risk: str
    rating: str                  # CRITICAL / HIGH / MEDIUM / LOW
    recommendation: str
    acceptance: str = ""
    source: str = ""             # finding / threat
    source_ref: str = ""         # finding fingerprint or "<run>/<threat id>"
    section: str = ""
    evidence: str = ""
    edited: bool = False
    # Design-specific wording drafted from the evidence; exported as the
    # Threat column when present (the generic rule sentence otherwise).
    threat_specific: str = ""
    recommendation_specific: str = ""


@dataclass
class PrelimRegister:
    document: str
    rows: List[PrelimRow] = field(default_factory=list)
    next_id: int = 1
    updated_at: str = ""

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: Dict[str, Any]) -> "PrelimRegister":
        return cls(rows=[PrelimRow(**r) for r in d.get("rows", [])],
                   **{k: v for k, v in d.items() if k != "rows"})


# --------------------------------------------------------------------------
# Building rows from confirmed findings and threats
# --------------------------------------------------------------------------
def _tokens(text: str) -> set:
    return {w for w in re.findall(r"[a-z0-9]+", (text or "").lower()) if len(w) > 3}


def _similar(a: str, b: str, threshold: float = 0.5) -> bool:
    x, y = _tokens(a), _tokens(b)
    if len(x) < 4 or len(y) < 4:       # too little text to call two rows the same
        return False
    return len(x & y) / len(x | y) >= threshold


def _threat_domain(dfd: Optional[DFD], t: Any) -> str:
    if dfd is None:
        return "security"
    cid = t.target_id
    if t.target_type == "flow":
        f = dfd.flow(t.target_id)
        cid = f.target if f else cid
    c = dfd.component(cid)
    return KIND_DOMAIN.get(c.kind, "application") if c else "security"


def candidates(findings: Iterable[Any], decisions: Dict[str, Dict[str, Any]],
               threat_runs: Iterable[Any] = (), dfd: Optional[DFD] = None,
               include: str = "accepted") -> List[PrelimRow]:
    """Confirmed items as rows (no IDs yet).

    include="accepted": only what the architect accepted.
    include="confirmed": also rule findings / rule threats not disputed.
    """
    rows: List[PrelimRow] = []
    for f in findings:
        d = decisions.get(f.fingerprint, {}).get("decision")
        if d == "disputed":
            continue
        if d != "accepted" and not (include == "confirmed" and f.origin == "rules_engine"):
            continue
        rows.append(PrelimRow(
            id="", domain=f.domain if f.domain in DEFAULT_DOMAINS else "security",
            threat=f.issue, risk="", rating=f.severity, recommendation=f.recommendation,
            source="finding", source_ref=f.fingerprint, section=f.section,
            evidence=f.evidence_excerpt))
    for run in threat_runs:
        for t in run.threats:
            if t.status in ("disputed", "mitigated"):
                continue
            if t.status != "accepted" and not (include == "confirmed" and t.source == "rule"):
                continue
            rows.append(PrelimRow(
                id="", domain=_threat_domain(dfd, t),
                threat=f"{t.title} ({t.framework} · {t.category}; {t.target_label})",
                risk="", rating=t.risk, recommendation="; ".join(t.mitigations),
                source="threat", source_ref=f"v{run.dfd_version}/{t.id}",
                section=f"DFD v{run.dfd_version}: {t.target_label}", evidence=t.description))
    # One row per problem: a threat and a finding about the same weakness merge,
    # keeping the higher rating.
    merged: List[PrelimRow] = []
    for r in sorted(rows, key=lambda x: ORDER.index(x.rating) if x.rating in ORDER else 9):
        if any(m.domain == r.domain and _similar(m.threat + " " + m.recommendation,
                                                 r.threat + " " + r.recommendation)
               for m in merged):
            continue
        merged.append(r)
    return merged


def sync_register(reg: PrelimRegister, fresh: List[PrelimRow]) -> List[str]:
    """Add new candidates with the next IDs; keep existing rows and their IDs.

    IDs never move or get reused: Stage 3 verifies against them.
    Returns the IDs added.
    """
    have = {r.source_ref for r in reg.rows}
    added = []
    for r in sorted(fresh, key=lambda x: (ORDER.index(x.rating) if x.rating in ORDER else 9,
                                           x.domain)):
        if r.source_ref in have:
            continue
        r.id = f"REC-{reg.next_id:03d}"
        reg.next_id += 1
        reg.rows.append(r)
        added.append(r.id)
    reg.updated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    return added


# --------------------------------------------------------------------------
# Risk statements and acceptance criteria (local model, architect edits)
# --------------------------------------------------------------------------
RISK_SYSTEM = """You write the "Threat", "Risk" and "Acceptance criteria" columns of a security architecture preliminary review register that goes to the delivery team.

For each numbered item you get the domain, a generic threat statement, the design section, the evidence quoted from the design, and the recommendation. Return ONLY JSON:
{"items": [{"n": <number>, "threat": <1 sentence>, "risk": <1-2 sentences>, "recommendation": <1-2 sentences>, "acceptance": <1-2 sentences>}]}

- threat: rewrite the generic threat for THIS design - name the actual component, route, port or setting from the evidence (e.g. "The /v1/payments/webhook route in Kong has authentication disabled"). Use only facts in the evidence.
- recommendation: rewrite the generic recommendation as the specific control for THIS component (keep the control intent, drop parts that do not apply - e.g. no CDN advice for a database).
- risk: the business consequence if the threat is realised - what an attacker achieves and what the organisation loses (data breach, fraud, outage, regulatory exposure). Do not restate the threat or give advice.
- acceptance: the concrete evidence an as-built review must see to consider the recommendation implemented (a configuration, setting, policy, diagram element or test result). Verifiable, specific, no "ensure" or "consider".
- Stay specific to the item. Compact JSON."""


def _message_text(response: Any) -> str:
    content = getattr(response, "content", response)
    if isinstance(content, list):
        return "".join(p.get("text", "") if isinstance(p, dict) else str(p) for p in content)
    return str(content or "")


def draft_risk_and_acceptance(rows: List[PrelimRow], llm: Any, batch_size: int = 6,
                              overwrite: bool = False) -> int:
    """Fill empty risk / acceptance fields; never touches an edited row. Returns rows filled."""
    todo = [r for r in rows if not r.edited and (overwrite or not r.risk or not r.acceptance
                                                 or not r.threat_specific)]
    filled = 0
    for i in range(0, len(todo), batch_size):
        chunk = todo[i:i + batch_size]
        user = "\n\n".join(
            f"{n}. DOMAIN: {DEFAULT_DOMAINS.get(r.domain, r.domain)}\nTHREAT: {r.threat}\n"
            f"SECTION: {r.section}\nEVIDENCE: {(r.evidence or '')[:400]}\n"
            f"RECOMMENDATION: {r.recommendation}" for n, r in enumerate(chunk, 1))
        try:
            raw = _message_text(llm.invoke([("system", RISK_SYSTEM), ("human", user)]))
            data = json.loads(re.search(r"\{.*\}", raw, re.S).group(0))
        except Exception:  # noqa: BLE001
            continue
        for item in data.get("items", []) if isinstance(data, dict) else []:
            try:
                r = chunk[int(item.get("n")) - 1]
            except (TypeError, ValueError, IndexError):
                continue
            specific = str(item.get("threat", "")).strip()
            if specific and (overwrite or not r.threat_specific):
                r.threat_specific = specific[:400]
            rec = str(item.get("recommendation", "")).strip()
            if rec and (overwrite or not r.recommendation_specific):
                r.recommendation_specific = rec[:600]
            if overwrite or not r.risk:
                r.risk = str(item.get("risk", "")).strip()[:600]
            if overwrite or not r.acceptance:
                r.acceptance = str(item.get("acceptance", "")).strip()[:600]
            filled += 1
    return filled


# --------------------------------------------------------------------------
# Storage and export
# --------------------------------------------------------------------------
def load(document: str, root: Path = DEFAULT_ROOT) -> PrelimRegister:
    p = root / slug(document) / "register.json"
    if not p.exists():
        return PrelimRegister(document=document)
    return PrelimRegister.from_dict(json.loads(p.read_text(encoding="utf-8")))


def save(reg: PrelimRegister, root: Path = DEFAULT_ROOT) -> Path:
    d = root / slug(reg.document)
    d.mkdir(parents=True, exist_ok=True)
    p = d / "register.json"
    p.write_text(json.dumps(reg.to_dict(), indent=2), encoding="utf-8")
    return p


def export_settings(config: Any = None) -> Dict[str, Any]:
    raw = (getattr(config, "raw", None) or {}).get("archer_export", {}) if config else {}
    return {
        "columns": {**DEFAULT_COLUMNS, **(raw.get("columns") or {})},
        "ratings": {**DEFAULT_RATINGS, **(raw.get("rating_values") or {})},
        "domains": {**DEFAULT_DOMAINS, **(raw.get("domain_values") or {})},
        "include_internal": bool(raw.get("include_id_and_acceptance", True)),
        "sheet": str(raw.get("sheet_name", "Prelim Review")),
    }


def table(reg: PrelimRegister, settings: Dict[str, Any]) -> List[Dict[str, str]]:
    cols = settings["columns"]
    keys = ["id", "domain", "threat", "risk", "rating", "recommendation", "acceptance"]
    if not settings["include_internal"]:
        keys = [k for k in keys if k not in ("id", "acceptance")]
    out = []
    for r in reg.rows:
        values = {"id": r.id, "domain": settings["domains"].get(r.domain, r.domain),
                  "threat": r.threat_specific or r.threat, "risk": r.risk,
                  "rating": settings["ratings"].get(r.rating, r.rating),
                  "recommendation": r.recommendation_specific or r.recommendation,
                  "acceptance": r.acceptance}
        out.append({cols[k]: values[k] for k in keys})
    return out


def to_xlsx(reg: PrelimRegister, settings: Dict[str, Any]) -> bytes:
    """Plain single-sheet workbook for Archer data import."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    rows = table(reg, settings)
    wb = Workbook()
    ws = wb.active
    ws.title = settings["sheet"][:31]
    headers = list(rows[0].keys()) if rows else list(table(
        PrelimRegister(document="", rows=[PrelimRow("", "", "", "", "", "")]), settings)[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row[h] for h in headers])
    widths = {"ID": 10, "Domain": 16, "Risk Rating": 12}
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(h, 48)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(reg: PrelimRegister, settings: Dict[str, Any]) -> str:
    import csv
    rows = table(reg, settings)
    buf = io.StringIO()
    if rows:
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return buf.getvalue()
