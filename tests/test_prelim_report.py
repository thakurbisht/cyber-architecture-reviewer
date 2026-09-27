"""Preliminary review register and Archer Excel export (src/prelim_report.py)."""

from __future__ import annotations

import io
import json

from src import dfd as D
from src import prelim_report as P
from src import threat_agent as TA
from src.models import Finding, ORIGIN_AGENT, ORIGIN_RULES


def _f(issue, severity="HIGH", domain="application", origin=ORIGIN_AGENT, section="API"):
    return Finding(section=section, domain=domain, severity=severity, issue=issue,
                   recommendation=f"Fix: {issue}", origin=origin)


def _run():
    d = D.DFD(document="x.md")
    d.components = [D.DFDComponent("db", "orders DB", kind="datastore", zone="restricted"),
                    D.DFDComponent("api", "API", kind="service", zone="dmz")]
    d.flows = [D.DFDFlow("f1", "api", "db", auth="password", encrypted="yes")]
    t_rule = TA.Threat(id="T-001", framework="STRIDE", category="Spoofing", target_type="flow",
                       target_id="f1", target_label="API → orders DB", title="Static password",
                       risk="HIGH", mitigations=["Use workload identity"], source="rule")
    t_llm = TA.Threat(id="T-002", framework="STRIDE", category="Tampering",
                      target_type="component", target_id="db", target_label="orders DB",
                      title="Backup tampering", risk="MEDIUM", mitigations=["Immutable backups"])
    return d, TA.ThreatModelRun(document="x.md", dfd_version=1, frameworks=["STRIDE"],
                                model="m", threats=[t_rule, t_llm])


def test_only_confirmed_items_become_rows():
    a = _f("Webhook route has no authentication", "CRITICAL")
    b = _f("Rate limit missing on cart", "MEDIUM")
    c = _f("SNMP v2c in use", "CRITICAL", "network", ORIGIN_RULES)
    decisions = {a.fingerprint: {"decision": "accepted"}, b.fingerprint: {"decision": "disputed"}}
    dfd, run = _run()
    accepted = P.candidates([a, b, c], decisions, [run], dfd, include="accepted")
    assert [r.threat for r in accepted] == ["Webhook route has no authentication"]
    confirmed = P.candidates([a, b, c], decisions, [run], dfd, include="confirmed")
    threats = {r.threat for r in confirmed}
    assert "SNMP v2c in use" in threats                          # rule finding
    assert any(t.startswith("Static password") for t in threats)  # rule threat
    assert not any(t.startswith("Backup tampering") for t in threats)  # not accepted
    assert "Rate limit missing on cart" not in threats          # disputed


def test_threat_domain_comes_from_dfd_element():
    dfd, run = _run()
    run.threats[1].status = "accepted"
    rows = P.candidates([], {}, [run], dfd, include="accepted")
    assert rows[0].domain == "cloud_data"                       # threat on a datastore


def test_duplicate_finding_and_threat_merge_keeping_higher_rating():
    a = _f("Orders DB uses a static password for the API", "MEDIUM", "cloud_data")
    a.recommendation = "Use workload identity instead of static password"
    dfd, run = _run()
    rows = P.candidates([a], {a.fingerprint: {"decision": "accepted"}}, [run], dfd,
                        include="confirmed")
    static = [r for r in rows if "static password" in r.threat.lower()]
    assert len(static) == 1 and static[0].rating == "HIGH"


def test_ids_are_stable_and_never_reused():
    reg = P.PrelimRegister(document="x.md")
    a = _f("Guest checkout creates accounts without email verification", "LOW")
    b = _f("Kong Admin API exposed without any authentication plugin", "CRITICAL")
    rows = P.candidates([a, b], {a.fingerprint: {"decision": "accepted"},
                                 b.fingerprint: {"decision": "accepted"}})
    assert P.sync_register(reg, rows) == ["REC-001", "REC-002"]
    assert reg.rows[0].threat.startswith("Kong Admin")                      # ordered by rating first time
    c = _f("Reporting replica publicly accessible from any address", "CRITICAL")
    again = P.candidates([a, b, c], {x.fingerprint: {"decision": "accepted"} for x in (a, b, c)})
    assert P.sync_register(reg, again) == ["REC-003"]           # existing rows keep their IDs
    assert [r.id for r in reg.rows] == ["REC-001", "REC-002", "REC-003"]


class FakeLLM:
    def invoke(self, messages):
        n = messages[-1][1].count("THREAT:")
        return json.dumps({"items": [{"n": i, "risk": f"risk {i}", "acceptance": f"acc {i}"}
                                     for i in range(1, n + 1)]})


def test_draft_fills_blanks_but_never_edited_rows():
    rows = [P.PrelimRow("REC-001", "network", "t1", "", "HIGH", "r1"),
            P.PrelimRow("REC-002", "network", "t2", "mine", "HIGH", "r2", edited=True)]
    assert P.draft_risk_and_acceptance(rows, FakeLLM()) == 1
    assert rows[0].risk == "risk 1" and rows[0].acceptance == "acc 1"
    assert rows[1].risk == "mine"


def test_excel_is_single_plain_sheet_with_configured_headers(tmp_path):
    from openpyxl import load_workbook
    reg = P.PrelimRegister(document="x.md", rows=[
        P.PrelimRow("REC-001", "cloud_data", "Public replica", "PII exposed", "CRITICAL",
                    "Make private", "publicly_accessible=false")])
    settings = P.export_settings()
    settings["columns"]["rating"] = "Inherent Risk Rating"      # Archer field name override
    wb = load_workbook(io.BytesIO(P.to_xlsx(reg, settings)))
    assert wb.sheetnames == ["Prelim Review"]
    ws = wb.active
    assert [c.value for c in ws[1]] == ["ID", "Domain", "Threat", "Risk", "Inherent Risk Rating",
                                        "Cyber Recommendation", "Acceptance Criteria"]
    assert [c.value for c in ws[2]][:5] == ["REC-001", "Cloud & Data", "Public replica",
                                            "PII exposed", "Critical"]
    assert not ws.merged_cells.ranges


def test_internal_columns_can_be_hidden_from_stakeholders():
    reg = P.PrelimRegister(document="x.md", rows=[P.PrelimRow("REC-001", "network", "t", "r",
                                                              "LOW", "rec", "acc")])
    s = P.export_settings()
    s["include_internal"] = False
    assert list(P.table(reg, s)[0].keys()) == ["Domain", "Threat", "Risk", "Risk Rating",
                                               "Cyber Recommendation"]


def test_register_round_trip(tmp_path):
    reg = P.PrelimRegister(document="x.md", rows=[P.PrelimRow("REC-001", "network", "t", "r",
                                                              "LOW", "rec")], next_id=2)
    P.save(reg, tmp_path)
    again = P.load("x.md", tmp_path)
    assert again.next_id == 2 and again.rows[0].id == "REC-001"
